import io
import pytest
import fitz
import docx
from app.services.parser import DocumentParserService, DocumentParserError


def create_sample_pdf(text: str) -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), text)
    pdf_bytes = doc.write()
    doc.close()
    return pdf_bytes


def create_sample_docx(text: str) -> bytes:
    doc = docx.Document()
    doc.add_heading("Resume", 0)
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_pdf_parsing_success():
    sample_text = (
        "Alex Developer\n"
        "alex@example.com | 555-0199 | San Francisco, CA\n\n"
        "Professional Summary:\n"
        "Senior Full Stack Engineer with 6 years of experience building React and FastAPI systems.\n\n"
        "Experience:\n"
        "Senior Software Engineer at CloudCorp (2021 - Present)\n"
        "- Scaled microservices architecture serving 2 million daily active users.\n"
        "- Reduced API latency by 45% using Redis caching and optimized queries.\n\n"
        "Education:\n"
        "BS Computer Science, University of California, Berkeley (2018)"
    )
    pdf_bytes = create_sample_pdf(sample_text)
    extracted, meta = DocumentParserService.parse_document(
        filename="resume.pdf",
        content=pdf_bytes,
        max_file_size_bytes=5 * 1024 * 1024
    )
    assert "Alex Developer" in extracted
    assert "CloudCorp" in extracted
    assert meta["format"] == "pdf"
    assert meta["pages"] >= 1


def test_docx_parsing_success():
    sample_text = (
        "Jordan Smith\n"
        "jordan@example.com | 555-0122 | Austin, TX\n\n"
        "Experienced Backend Developer specializing in Python, FastAPI, and Postgres.\n"
        "Built robust payment integrations and high-throughput background processing."
    )
    docx_bytes = create_sample_docx(sample_text)
    extracted, meta = DocumentParserService.parse_document(
        filename="resume.docx",
        content=docx_bytes,
        max_file_size_bytes=5 * 1024 * 1024
    )
    assert "Jordan Smith" in extracted
    assert "FastAPI" in extracted
    assert meta["format"] == "docx"


def test_invalid_extension_or_format():
    with pytest.raises(DocumentParserError) as exc_info:
        DocumentParserService.parse_document("test.txt", b"plain text", 1024 * 1024)
    assert exc_info.value.status_code == 415

    # File named .pdf but containing plain text (invalid magic bytes)
    with pytest.raises(DocumentParserError) as exc_info2:
        DocumentParserService.parse_document("fake.pdf", b"this is not a pdf file", 1024 * 1024)
    assert exc_info2.value.status_code == 415


def test_empty_file():
    with pytest.raises(DocumentParserError) as exc_info:
        DocumentParserService.parse_document("empty.pdf", b"", 1024 * 1024)
    assert exc_info.value.status_code == 400


def test_oversized_file():
    with pytest.raises(DocumentParserError) as exc_info:
        DocumentParserService.parse_document("large.pdf", b"%PDF-" + b"0" * 1000, 500)
    assert exc_info.value.status_code == 413


def test_empty_pdf_detection():
    # PDF with blank page and no text
    doc = fitz.open()
    doc.new_page()
    pdf_bytes = doc.write()
    doc.close()

    with pytest.raises(DocumentParserError) as exc_info:
        DocumentParserService.parse_document("blank.pdf", pdf_bytes, 1024 * 1024)
    assert exc_info.value.status_code == 422
    assert "No readable text" in str(exc_info.value.message)
