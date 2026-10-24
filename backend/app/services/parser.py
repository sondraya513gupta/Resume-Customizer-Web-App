import io
import re
from typing import Tuple, Dict, Any
from fastapi import HTTPException, status
import pymupdf as fitz  # PyMuPDF
import docx  # python-docx


class DocumentParserError(Exception):
    """Custom exception for document parsing failures."""
    def __init__(self, message: str, status_code: int = 422):
        super().__init__(message)
        self.message = message
        self.status_code = status_code


class DocumentParserService:
    """Service for validating, reading, and extracting text from PDF and DOCX files."""

    PDF_MAGIC = b"%PDF-"
    DOCX_MAGIC = b"PK\x03\x04"
    MIN_CHARACTERS_REQUIRED = 50

    @classmethod
    def validate_file_format(cls, filename: str, content: bytes) -> str:
        """
        Validate file extension and magic byte headers.
        Returns normalized file type: 'pdf' or 'docx'.
        """
        if not filename:
            raise DocumentParserError("Filename is missing.", status.HTTP_400_BAD_REQUEST)

        filename_lower = filename.lower()
        if filename_lower.endswith(".pdf"):
            if not content.startswith(cls.PDF_MAGIC):
                raise DocumentParserError(
                    "File has a .pdf extension but does not contain a valid PDF signature.",
                    status.HTTP_415_UNSUPPORTED_MEDIA_TYPE
                )
            return "pdf"
        elif filename_lower.endswith(".docx"):
            if not content.startswith(cls.DOCX_MAGIC):
                raise DocumentParserError(
                    "File has a .docx extension but does not contain a valid DOCX signature.",
                    status.HTTP_415_UNSUPPORTED_MEDIA_TYPE
                )
            return "docx"
        elif filename_lower.endswith(".doc"):
            raise DocumentParserError(
                "Legacy .doc files are not supported. Please convert your file to .docx or .pdf before uploading.",
                status.HTTP_415_UNSUPPORTED_MEDIA_TYPE
            )
        else:
            raise DocumentParserError(
                "Unsupported file type. Only PDF (.pdf) and Microsoft Word (.docx) files are supported.",
                status.HTTP_415_UNSUPPORTED_MEDIA_TYPE
            )

    @classmethod
    def extract_text_from_pdf(cls, content: bytes) -> Tuple[str, Dict[str, Any]]:
        """
        Extract text from PDF using PyMuPDF (fitz).
        Detects encrypted files, corrupted files, and scanned image-only PDFs.
        """
        try:
            doc = fitz.open(stream=content, filetype="pdf")
        except Exception as e:
            raise DocumentParserError(f"Failed to open PDF document. The file may be corrupted: {str(e)}")

        try:
            if doc.is_encrypted:
                raise DocumentParserError(
                    "The uploaded PDF is password-protected or encrypted. Please upload an unencrypted document."
                )

            page_count = len(doc)
            if page_count == 0:
                raise DocumentParserError("The uploaded PDF has no pages.")

            extracted_pages = []
            total_images = 0

            for page_num in range(page_count):
                page = doc[page_num]
                text = page.get_text("text") or ""
                extracted_pages.append(text)
                total_images += len(page.get_images())

            full_text = "\n\n".join(extracted_pages).strip()
            # Clean up null characters or unusual control bytes
            full_text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", full_text)

            # Check if text is essentially empty or scanned
            clean_word_count = len(full_text.split())
            if len(full_text) < cls.MIN_CHARACTERS_REQUIRED or clean_word_count < 10:
                if total_images > 0:
                    raise DocumentParserError(
                        "The uploaded PDF appears to be a scanned image or photograph. ResumeCraft AI requires "
                        "a digital text-based resume (OCR is not supported). Please upload a selectable-text PDF or DOCX."
                    )
                else:
                    raise DocumentParserError(
                        "No readable text could be extracted from this PDF. Please ensure the document is not empty."
                    )

            metadata = {
                "format": "pdf",
                "pages": page_count,
                "character_count": len(full_text),
                "word_count": clean_word_count,
            }
            return full_text, metadata
        finally:
            doc.close()

    @classmethod
    def extract_text_from_docx(cls, content: bytes) -> Tuple[str, Dict[str, Any]]:
        """
        Extract text from DOCX using python-docx.
        Extracts both paragraphs and table cells.
        """
        try:
            doc = docx.Document(io.BytesIO(content))
        except Exception as e:
            raise DocumentParserError(f"Failed to open DOCX document. The file may be corrupted: {str(e)}")

        paragraphs_text = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                paragraphs_text.append(text)

        # Also extract table text as many resumes use tables for layouts
        table_text = []
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    table_text.append(" | ".join(row_cells))

        combined_text = "\n\n".join(paragraphs_text + table_text).strip()
        combined_text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", combined_text)
        clean_word_count = len(combined_text.split())

        if len(combined_text) < cls.MIN_CHARACTERS_REQUIRED or clean_word_count < 10:
            raise DocumentParserError(
                "The uploaded DOCX file does not contain sufficient extractable text. "
                "Please verify the file contains your resume text and is not empty."
            )

        metadata = {
            "format": "docx",
            "paragraphs": len(paragraphs_text),
            "tables": len(doc.tables),
            "character_count": len(combined_text),
            "word_count": clean_word_count,
        }
        return combined_text, metadata

    @classmethod
    def parse_document(cls, filename: str, content: bytes, max_file_size_bytes: int) -> Tuple[str, Dict[str, Any]]:
        """
        Comprehensive pipeline: validates size, format, extracts text, and returns metadata.
        """
        if len(content) == 0:
            raise DocumentParserError("The uploaded file is empty (0 bytes).", status.HTTP_400_BAD_REQUEST)

        if len(content) > max_file_size_bytes:
            max_mb = max_file_size_bytes / (1024 * 1024)
            raise DocumentParserError(
                f"File size exceeds the {max_mb:.0f} MB limit. Please upload a smaller file.",
                status.HTTP_413_REQUEST_ENTITY_TOO_LARGE
            )

        fmt = cls.validate_file_format(filename, content)
        if fmt == "pdf":
            return cls.extract_text_from_pdf(content)
        elif fmt == "docx":
            return cls.extract_text_from_docx(content)
        else:
            raise DocumentParserError("Unsupported format.", status.HTTP_415_UNSUPPORTED_MEDIA_TYPE)
