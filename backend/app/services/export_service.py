import io
from typing import Optional
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.schemas.resume import CustomizedResume


class ExportService:
    """Generates production-grade ATS-friendly PDF and DOCX documents."""

    @classmethod
    def generate_pdf(cls, resume: CustomizedResume) -> bytes:
        """
        Builds a crisp, professional, ATS-optimized PDF resume using ReportLab.
        Ensures selectable text, clean margins, and proper hierarchy.
        """
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            leftMargin=36,  # 0.5 inch
            rightMargin=36,
            topMargin=36,
            bottomMargin=36
        )

        styles = getSampleStyleSheet()

        # Custom typography styles
        name_style = ParagraphStyle(
            "CandidateName",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=18,
            leading=22,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#0f172a")  # Slate 900
        )

        contact_style = ParagraphStyle(
            "ContactInfo",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=9,
            leading=12,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#475569")  # Slate 600
        )

        section_heading = ParagraphStyle(
            "SectionHeading",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=11,
            leading=14,
            textColor=colors.HexColor("#0f172a"),
            spaceBefore=8,
            spaceAfter=4,
            keepWithNext=True
        )

        body_style = ParagraphStyle(
            "ResumeBody",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=9.5,
            leading=13,
            textColor=colors.HexColor("#1e293b")
        )

        bold_style = ParagraphStyle(
            "ResumeBold",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=9.5,
            leading=13,
            textColor=colors.HexColor("#0f172a")
        )

        bullet_style = ParagraphStyle(
            "ResumeBullet",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=9,
            leading=12.5,
            leftIndent=14,
            firstLineIndent=-10,
            textColor=colors.HexColor("#1e293b"),
            spaceAfter=3
        )

        subtext_style = ParagraphStyle(
            "ResumeSubtext",
            parent=styles["Normal"],
            fontName="Helvetica-Oblique",
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor("#64748b")
        )

        story = []

        # 1. Header: Name & Contact Info
        name = resume.candidate_name or "Candidate Name"
        story.append(Paragraph(name.upper(), name_style))
        story.append(Spacer(1, 4))

        contact_parts = []
        if resume.contact_info.get("email"):
            contact_parts.append(resume.contact_info["email"])
        if resume.contact_info.get("phone"):
            contact_parts.append(resume.contact_info["phone"])
        if resume.contact_info.get("location"):
            contact_parts.append(resume.contact_info["location"])
        if resume.contact_info.get("linkedin"):
            contact_parts.append(resume.contact_info["linkedin"])
        if resume.contact_info.get("github"):
            contact_parts.append(resume.contact_info["github"])

        if contact_parts:
            contact_line = "  •  ".join(contact_parts)
            story.append(Paragraph(contact_line, contact_style))
            story.append(Spacer(1, 8))

        # Divider
        story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#cbd5e1"), spaceBefore=2, spaceAfter=8))

        # 2. Professional Summary
        if resume.tailored_summary:
            story.append(Paragraph("PROFESSIONAL SUMMARY", section_heading))
            story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#94a3b8"), spaceBefore=1, spaceAfter=4))
            story.append(Paragraph(resume.tailored_summary, body_style))
            story.append(Spacer(1, 6))

        # 3. Technical & Professional Skills
        if resume.skills:
            story.append(Paragraph("SKILLS & COMPETENCIES", section_heading))
            story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#94a3b8"), spaceBefore=1, spaceAfter=4))
            for cat in resume.skills:
                skills_str = ", ".join(cat.skills)
                cat_paragraph = Paragraph(
                    f"<b>{cat.category}:</b> {skills_str}",
                    body_style
                )
                story.append(cat_paragraph)
            story.append(Spacer(1, 6))

        # 4. Experience
        if resume.experience:
            story.append(Paragraph("PROFESSIONAL EXPERIENCE", section_heading))
            story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#94a3b8"), spaceBefore=1, spaceAfter=4))

            for exp in resume.experience:
                # Top row: Role & Company (left) + Dates (right)
                left_text = f"<b>{exp.role}</b> — <i>{exp.company}</i>"
                if exp.location:
                    left_text += f" ({exp.location})"

                date_text = ""
                if exp.start_date or exp.end_date:
                    date_text = f"{exp.start_date or ''} - {exp.end_date or 'Present'}"

                table_data = [
                    [Paragraph(left_text, body_style), Paragraph(date_text, ParagraphStyle("RightDate", parent=subtext_style, alignment=TA_RIGHT))]
                ]
                row_table = Table(table_data, colWidths=[400, 140])
                row_table.setStyle(TableStyle([
                    ('VALIGN', (0,0), (-1,-1), 'TOP'),
                    ('LEFTPADDING', (0,0), (-1,-1), 0),
                    ('RIGHTPADDING', (0,0), (-1,-1), 0),
                    ('TOPPADDING', (0,0), (-1,-1), 2),
                    ('BOTTOMPADDING', (0,0), (-1,-1), 2),
                ]))
                story.append(row_table)

                # Bullets
                bullets = exp.tailored_bullets or exp.original_bullets
                for bullet in bullets:
                    bullet_text = f"&bull; {bullet}"
                    story.append(Paragraph(bullet_text, bullet_style))
                story.append(Spacer(1, 4))

        # 5. Projects
        if resume.projects:
            story.append(Paragraph("TECHNICAL PROJECTS", section_heading))
            story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#94a3b8"), spaceBefore=1, spaceAfter=4))
            for proj in resume.projects:
                tech_str = f" | <i>{', '.join(proj.technologies)}</i>" if proj.technologies else ""
                proj_title = f"<b>{proj.name}</b>{tech_str}"
                story.append(Paragraph(proj_title, body_style))
                story.append(Paragraph(f"&bull; {proj.description}", bullet_style))
                story.append(Spacer(1, 3))

        # 6. Education
        if resume.education:
            story.append(Paragraph("EDUCATION", section_heading))
            story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#94a3b8"), spaceBefore=1, spaceAfter=4))
            for edu in resume.education:
                deg_field = edu.degree
                if edu.field_of_study:
                    deg_field += f" in {edu.field_of_study}"
                left_text = f"<b>{deg_field}</b> — {edu.institution}"
                date_text = edu.graduation_date or ""
                table_data = [
                    [Paragraph(left_text, body_style), Paragraph(date_text, ParagraphStyle("RightEduDate", parent=subtext_style, alignment=TA_RIGHT))]
                ]
                row_table = Table(table_data, colWidths=[420, 120])
                row_table.setStyle(TableStyle([
                    ('VALIGN', (0,0), (-1,-1), 'TOP'),
                    ('LEFTPADDING', (0,0), (-1,-1), 0),
                    ('RIGHTPADDING', (0,0), (-1,-1), 0),
                    ('TOPPADDING', (0,0), (-1,-1), 1),
                    ('BOTTOMPADDING', (0,0), (-1,-1), 1),
                ]))
                story.append(row_table)
            story.append(Spacer(1, 4))

        # 7. Certifications
        if resume.certifications:
            story.append(Paragraph("CERTIFICATIONS", section_heading))
            story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#94a3b8"), spaceBefore=1, spaceAfter=4))
            for cert in resume.certifications:
                story.append(Paragraph(f"&bull; {cert}", bullet_style))

        doc.build(story)
        buffer.seek(0)
        return buffer.getvalue()

    @classmethod
    def generate_docx(cls, resume: CustomizedResume) -> bytes:
        """Generates a downloadable DOCX document matching the tailored resume."""
        doc = docx.Document()

        # Set 0.6 inch margins
        for section in doc.sections:
            section.top_margin = Inches(0.6)
            section.bottom_margin = Inches(0.6)
            section.left_margin = Inches(0.6)
            section.right_margin = Inches(0.6)

        # Name Heading
        name_p = doc.add_paragraph()
        name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = name_p.add_run(resume.candidate_name.upper())
        run.bold = True
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor(15, 23, 42)

        # Contact line
        contact_parts = []
        for k in ["email", "phone", "location", "linkedin", "github"]:
            if resume.contact_info.get(k):
                contact_parts.append(resume.contact_info[k])
        if contact_parts:
            cp = doc.add_paragraph(" | ".join(contact_parts))
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.space_after = Pt(12)

        # Helper for section headings
        def add_heading(text: str):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(3)
            r = h.add_run(text)
            r.bold = True
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(15, 23, 42)

        # Summary
        if resume.tailored_summary:
            add_heading("PROFESSIONAL SUMMARY")
            sp = doc.add_paragraph(resume.tailored_summary)
            sp.paragraph_format.space_after = Pt(6)

        # Skills
        if resume.skills:
            add_heading("SKILLS & COMPETENCIES")
            for cat in resume.skills:
                kp = doc.add_paragraph()
                r_cat = kp.add_run(f"{cat.category}: ")
                r_cat.bold = True
                kp.add_run(", ".join(cat.skills))
                kp.paragraph_format.space_after = Pt(2)

        # Experience
        if resume.experience:
            add_heading("PROFESSIONAL EXPERIENCE")
            for exp in resume.experience:
                ep = doc.add_paragraph()
                r_title = ep.add_run(f"{exp.role} — {exp.company}")
                r_title.bold = True
                dates = f" ({exp.start_date or ''} - {exp.end_date or 'Present'})"
                ep.add_run(dates)
                ep.paragraph_format.space_after = Pt(2)

                bullets = exp.tailored_bullets or exp.original_bullets
                for b in bullets:
                    bp = doc.add_paragraph(b, style="List Bullet")
                    bp.paragraph_format.space_after = Pt(2)

        # Education
        if resume.education:
            add_heading("EDUCATION")
            for edu in resume.education:
                edup = doc.add_paragraph()
                deg_str = f"{edu.degree}" + (f" in {edu.field_of_study}" if edu.field_of_study else "")
                r_deg = edup.add_run(f"{deg_str} — {edu.institution}")
                r_deg.bold = True
                if edu.graduation_date:
                    edup.add_run(f" ({edu.graduation_date})")
                edup.paragraph_format.space_after = Pt(2)

        # Projects
        if resume.projects:
            add_heading("PROJECTS")
            for proj in resume.projects:
                pp = doc.add_paragraph()
                r_p = pp.add_run(proj.name)
                r_p.bold = True
                if proj.technologies:
                    pp.add_run(f" [{', '.join(proj.technologies)}]")
                pp.paragraph_format.space_after = Pt(2)
                bp = doc.add_paragraph(proj.description, style="List Bullet")
                bp.paragraph_format.space_after = Pt(2)

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    @classmethod
    def generate_cover_letter_pdf(cls, resume: CustomizedResume) -> bytes:
        """Generates a professional PDF for the generated cover letter."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            leftMargin=54,
            rightMargin=54,
            topMargin=54,
            bottomMargin=54
        )

        styles = getSampleStyleSheet()

        name_style = ParagraphStyle(
            "CLName",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=16,
            textColor=colors.HexColor("#0f172a")
        )

        contact_style = ParagraphStyle(
            "CLContact",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=9,
            textColor=colors.HexColor("#475569")
        )

        body_style = ParagraphStyle(
            "CLBody",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=10.5,
            leading=16,
            textColor=colors.HexColor("#1e293b"),
            spaceAfter=12
        )

        story = []
        story.append(Paragraph(resume.candidate_name, name_style))

        contact_parts = []
        for k in ["email", "phone", "location", "linkedin"]:
            if resume.contact_info.get(k):
                contact_parts.append(resume.contact_info[k])
        if contact_parts:
            story.append(Paragraph(" | ".join(contact_parts), contact_style))

        story.append(Spacer(1, 16))
        story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#cbd5e1"), spaceAfter=18))

        # Add cover letter paragraphs
        cl_text = resume.cover_letter or "No cover letter was generated."
        paragraphs = [p.strip() for p in cl_text.split("\n\n") if p.strip()]

        for p in paragraphs:
            story.append(Paragraph(p.replace("\n", "<br/>"), body_style))

        doc.build(story)
        buffer.seek(0)
        return buffer.getvalue()
